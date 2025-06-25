#!/miniconda3/bin/python
#-*- coding: utf-8 -*-
#File Name    : xlsx2word.py
#Author       : liangjifeng
#Mail         : jifeng.liang@biobin.com.cn
#Create Time  : 2025-06-10 11:27

#!/usr/bin/env python3
from docx import Document
import pandas as pd

def convert_excel_to_word(excel_path, word_path):
    doc = Document()
    df = pd.read_excel(excel_path)
    
    # 添加表格到Word
    table = doc.add_table(df.shape[0]+1, df.shape[1])
    for col_idx, col_name in enumerate(df.columns):
        table.cell(0, col_idx).text = str(col_name)
    
    for row_idx, row in df.iterrows():
        for col_idx, value in enumerate(row):
            table.cell(row_idx+1, col_idx).text = str(value)
    
    doc.save(word_path)

if __name__ == "__main__":
    convert_excel_to_word("ANOVA.xlsx", "output.docx")
